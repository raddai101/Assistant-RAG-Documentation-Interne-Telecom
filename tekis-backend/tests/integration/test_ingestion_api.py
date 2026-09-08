import io

from app import create_app
from app.extensions import db


def _make_app(tmp_path):
    app = create_app("testing")
    app.config["INGESTION_STORAGE_DIR"] = str(tmp_path / "documents")
    app.config["QUARANTINE_DIR"] = str(tmp_path / "quarantine")
    with app.app_context():
        db.create_all()
    return app


def _upload(client, filename, content_bytes, title, **extra_fields):
    data = {"title": title, "file": (io.BytesIO(content_bytes), filename)}
    data.update(extra_fields)
    return client.post("/api/v1/ingestion", data=data, content_type="multipart/form-data")


def test_ingest_txt_document_creates_document_and_chunks(tmp_path):
    app = _make_app(tmp_path)
    client = app.test_client()
    content = ("Procédure réseau SGSN. " * 100).encode("utf-8")

    response = _upload(client, "procedure.txt", content, "Procédure SGSN")

    assert response.status_code == 201
    body = response.get_json()
    assert body["success"] is True
    assert body["data"]["version_number"] == 1
    assert body["data"]["num_chunks"] > 0

    document_id = body["data"]["document_id"]
    get_response = client.get(f"/api/v1/documents/{document_id}")
    assert get_response.status_code == 200
    doc_body = get_response.get_json()
    assert doc_body["data"]["title"] == "Procédure SGSN"
    assert len(doc_body["data"]["versions"]) == 1
    assert doc_body["data"]["versions"][0]["status"] == "active"


def test_ingest_missing_title_returns_400(tmp_path):
    app = _make_app(tmp_path)
    client = app.test_client()

    response = client.post(
        "/api/v1/ingestion",
        data={"file": (io.BytesIO(b"contenu"), "note.txt")},
        content_type="multipart/form-data",
    )

    assert response.status_code == 400
    assert response.get_json()["success"] is False


def test_ingest_unsupported_extension_returns_400(tmp_path):
    app = _make_app(tmp_path)
    client = app.test_client()

    response = _upload(client, "schema.dwg", b"binaire", "Schéma antenne")

    assert response.status_code == 400
    assert "non supportée" in response.get_json()["error"]


def test_ingest_new_version_of_existing_document_supersedes_previous(tmp_path):
    app = _make_app(tmp_path)
    client = app.test_client()

    first = _upload(client, "v1.txt", ("Version 1 du document. " * 50).encode(), "Norme antenne")
    document_id = first.get_json()["data"]["document_id"]
    first_version_id = first.get_json()["data"]["document_version_id"]

    second = _upload(
        client,
        "v2.txt",
        ("Version 2 corrigée du document. " * 50).encode(),
        "Norme antenne",
        document_id=str(document_id),
        supersedes_version_id=str(first_version_id),
    )

    assert second.status_code == 201
    assert second.get_json()["data"]["version_number"] == 2

    doc_body = client.get(f"/api/v1/documents/{document_id}").get_json()
    versions = {v["id"]: v for v in doc_body["data"]["versions"]}
    assert versions[first_version_id]["status"] == "superseded"
    assert len(doc_body["data"]["versions"]) == 2


def test_ingest_txt_content_disguised_as_pdf_is_rejected_and_quarantined(tmp_path):
    """Vérifie la détection de type réel (MIME) : contenu texte brut renommé en
    .pdf doit être rejeté malgré l'extension déclarée."""
    app = _make_app(tmp_path)
    client = app.test_client()

    response = _upload(client, "faux.pdf", b"Ceci n'est pas un PDF, juste du texte.", "Document suspect")

    assert response.status_code == 400
    body = response.get_json()
    assert body["success"] is False
    assert "PDF" in body["error"]

    # Le fichier doit être en quarantaine/rejected, jamais dans le stockage définitif
    # ni supprimé silencieusement.
    quarantine_dir = tmp_path / "quarantine"
    storage_dir = tmp_path / "documents"
    rejected_files = list((quarantine_dir / "rejected").glob("*.pdf"))
    reason_files = list((quarantine_dir / "rejected").glob("*.reason.txt"))
    assert len(rejected_files) == 1
    assert len(reason_files) == 1
    assert not storage_dir.exists() or not any(storage_dir.iterdir())


def test_ingest_valid_docx_and_xlsx_pass_mime_validation(tmp_path):
    import docx
    import openpyxl

    app = _make_app(tmp_path)
    client = app.test_client()

    docx_path = tmp_path / "note.docx"
    document = docx.Document()
    document.add_paragraph("Note de service télécom valide.")
    document.save(docx_path)
    with open(docx_path, "rb") as f:
        docx_bytes = f.read()

    response = _upload(client, "note.docx", docx_bytes, "Note de service")
    assert response.status_code == 201

    xlsx_path = tmp_path / "inventaire.xlsx"
    workbook = openpyxl.Workbook()
    workbook.active.append(["Site", "Statut"])
    workbook.save(xlsx_path)
    with open(xlsx_path, "rb") as f:
        xlsx_bytes = f.read()

    response2 = _upload(client, "inventaire.xlsx", xlsx_bytes, "Inventaire équipements")
    assert response2.status_code == 201


def test_ingest_empty_file_is_rejected(tmp_path):
    app = _make_app(tmp_path)
    client = app.test_client()

    response = _upload(client, "vide.txt", b"", "Fichier vide")

    assert response.status_code == 400
    assert "vide" in response.get_json()["error"].lower()


def test_ingest_oversized_file_returns_413(tmp_path):
    app = _make_app(tmp_path)
    app.config["MAX_UPLOAD_SIZE_MB"] = 1
    app.config["MAX_CONTENT_LENGTH"] = 1 * 1024 * 1024  # cohérent avec MAX_UPLOAD_SIZE_MB
    client = app.test_client()

    oversized_content = b"A" * (2 * 1024 * 1024)  # 2 Mo > limite de 1 Mo

    response = _upload(client, "gros.txt", oversized_content, "Fichier trop gros")

    assert response.status_code == 413


def test_ingest_pdf_with_valid_signature_but_corrupted_content_returns_400(tmp_path):
    """Régression : un PDF avec un en-tête %PDF- valide mais un contenu tronqué
    passait la validation de signature puis faisait planter le parsing avec une
    erreur 500 non gérée. Doit maintenant être rejeté proprement en 400."""
    app = _make_app(tmp_path)
    client = app.test_client()

    truncated_pdf = b"%PDF-1.4\n%contenu tronque sans structure PDF valide derriere"

    response = _upload(client, "corrompu.pdf", truncated_pdf, "PDF corrompu")

    assert response.status_code == 400
    body = response.get_json()
    assert body["success"] is False
    assert "corrompu" in body["error"].lower() or "illisible" in body["error"].lower()

    rejected_files = list((tmp_path / "quarantine" / "rejected").glob("*.pdf"))
    assert len(rejected_files) == 1


def test_successful_ingestion_leaves_quarantine_empty(tmp_path):
    app = _make_app(tmp_path)
    client = app.test_client()

    response = _upload(client, "propre.txt", ("Contenu propre et valide. " * 30).encode(), "Doc propre")

    assert response.status_code == 201
    quarantine_dir = tmp_path / "quarantine"
    remaining = [p for p in quarantine_dir.iterdir() if p.is_file()]
    assert remaining == []
    storage_dir = tmp_path / "documents"
    assert len(list(storage_dir.iterdir())) == 1
