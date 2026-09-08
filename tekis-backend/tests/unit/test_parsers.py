import os
import tempfile

import docx
import openpyxl
import pytest

from app.modules.ingestion.parsers import get_parser_for, UnsupportedFileTypeError
from app.modules.ingestion.parsers.txt_parser import TxtParser
from app.modules.ingestion.parsers.docx_parser import DocxParser
from app.modules.ingestion.parsers.xlsx_parser import XlsxParser


def test_txt_parser_reads_content():
    with tempfile.NamedTemporaryFile(mode="w", suffix=".txt", delete=False) as f:
        f.write("Procédure SGSN-2024-17.\nContenu de test.")
        path = f.name

    try:
        parsed = TxtParser().parse(path)
        assert "SGSN-2024-17" in parsed.text
        assert parsed.metadata["format"] == "txt"
    finally:
        os.unlink(path)


def test_docx_parser_reads_paragraphs():
    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as f:
        path = f.name
    document = docx.Document()
    document.add_paragraph("Titre de la procédure")
    document.add_paragraph("Corps du document télécom.")
    document.save(path)

    try:
        parsed = DocxParser().parse(path)
        assert "Titre de la procédure" in parsed.text
        assert "Corps du document télécom." in parsed.text
        assert parsed.metadata["num_paragraphs"] == 2
    finally:
        os.unlink(path)


def test_xlsx_parser_reads_sheets():
    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
        path = f.name
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Equipements"
    sheet.append(["Site", "Type", "Statut"])
    sheet.append(["Kinshasa-BTS-01", "BTS", "actif"])
    workbook.save(path)

    try:
        parsed = XlsxParser().parse(path)
        assert "Kinshasa-BTS-01" in parsed.text
        assert "[Feuille: Equipements]" in parsed.text
        assert parsed.metadata["num_sheets"] == 1
    finally:
        os.unlink(path)


def test_get_parser_for_unsupported_extension_raises():
    with pytest.raises(UnsupportedFileTypeError):
        get_parser_for("fichier.xyz")


def test_get_parser_for_known_extensions():
    assert get_parser_for("doc.pdf").supported_extensions == (".pdf",)
    assert get_parser_for("doc.docx").supported_extensions == (".docx",)
    assert get_parser_for("doc.xlsx").supported_extensions == (".xlsx",)
    assert get_parser_for("doc.txt").supported_extensions == (".txt",)
