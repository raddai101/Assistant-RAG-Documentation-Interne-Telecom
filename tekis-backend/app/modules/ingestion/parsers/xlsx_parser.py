import openpyxl

from app.modules.ingestion.contracts import ParsedDocument


class XlsxParser:
    supported_extensions = (".xlsx",)

    def parse(self, file_path: str) -> ParsedDocument:
        workbook = openpyxl.load_workbook(file_path, data_only=True, read_only=True)
        try:
            sheet_texts = []
            for sheet in workbook.worksheets:
                rows_text = []
                for row in sheet.iter_rows(values_only=True):
                    cells = [str(c) for c in row if c is not None]
                    if cells:
                        rows_text.append(" | ".join(cells))
                sheet_text = f"[Feuille: {sheet.title}]\n" + "\n".join(rows_text)
                sheet_texts.append(sheet_text)

            full_text = "\n\n".join(sheet_texts)
            metadata = {"format": "xlsx", "num_sheets": len(workbook.worksheets)}
            return ParsedDocument(text=full_text, pages=sheet_texts, metadata=metadata)
        finally:
            workbook.close()
