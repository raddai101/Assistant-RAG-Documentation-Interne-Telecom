import docx

from app.modules.ingestion.contracts import ParsedDocument


class DocxParser:
    supported_extensions = (".docx",)

    def parse(self, file_path: str) -> ParsedDocument:
        document = docx.Document(file_path)
        paragraphs = [p.text for p in document.paragraphs if p.text.strip()]
        full_text = "\n".join(paragraphs)
        metadata = {"format": "docx", "num_paragraphs": len(paragraphs)}
        # Pas de notion native de "page" en docx : on traite le document comme une
        # seule page logique, la segmentation fine se fait au chunking.
        return ParsedDocument(text=full_text, pages=[full_text], metadata=metadata)
